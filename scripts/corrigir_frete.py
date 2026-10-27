"""D-067: ferramenta separada, NUNCA chamada pela rotina ou pelo Actions.

Simulação padrão; aplicar exige --aplicar E a frase de autorização explícita.
Não foi executada contra a loja durante a implementação.
"""
import argparse
from datetime import datetime
import json
from pathlib import Path
import re
import sys
from urllib.parse import urlsplit

from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from painel_monstro.log_publico import comando_seguro, ArgumentParserSeguro, resumo_erro
from painel_monstro.config import carregar
from painel_monstro.coletor.nuvemshop_api import NuvemshopAPI
from painel_monstro.saidas.cadastro import numero

FRASE='pode aplicar as correções de frete'
CAMPOS={'weight':'Peso proposto (kg)','height':'Altura proposta (cm)',
        'width':'Largura proposta (cm)','depth':'Comprimento proposto (cm)'}


def ler_propostas(caminho):
    wb=load_workbook(caminho,read_only=True,data_only=True)
    linhas=[]
    try:
        for ws in wb:
            rs=ws.iter_rows(values_only=True); cabe=next(rs,())
            if not any(str(c).startswith('Autoriza') for c in cabe): continue
            for valores in rs:
                r=dict(zip(cabe,valores))
                autoriza=next((v for k,v in r.items() if str(k).startswith('Autoriza')),'')
                if str(autoriza).strip().lower()!='sim': continue
                proposta={k:numero(r.get(c)) for k,c in CAMPOS.items()}
                if not all(proposta.values()): raise ValueError('Linha autorizada sem peso ou medidas positivos.')
                if any(proposta[k]<1 for k in ('height','width','depth')):
                    raise ValueError('Medidas devem estar em centímetros, nunca 0,32.')
                # Resolve ID explícito ou link da loja; nunca inferir pelo nome do produto.
                pid=str(r.get('produto_id') or r.get('Produto ID') or '')
                if not pid:
                    m=re.search(r'/admin/products/(\d+)(?:[/?#]|$)',str(r.get('Link','')))
                    pid=m[1] if m else ''
                handle=''
                if not pid:
                    u=urlsplit(str(r.get('Link','')))
                    m=re.fullmatch(r'/produtos/([\w-]+)/?',u.path)
                    if u.hostname=='lojamonstro.com.br' and m: handle=m[1]
                if not pid.isdigit() and not handle: raise ValueError('Linha autorizada sem ID ou link válido da loja.')
                linhas.append({'produto_id':pid,'handle':handle,'proposta':proposta})
    finally: wb.close()
    if len({(r['produto_id'],r['handle']) for r in linhas})!=len(linhas): raise ValueError('Produto autorizado duplicado na planilha.')
    return linhas


def executar(api,propostas,relatorio,aplicar=False,autorizacao=''):
    if aplicar and autorizacao!=FRASE: raise ValueError('Aplicação bloqueada: falta autorização escrita exata.')
    caminho=Path(relatorio).resolve()
    raiz=(Path(__file__).resolve().parents[1]/'dados').resolve()
    if not caminho.is_relative_to(raiz): raise ValueError('Relatório deve ficar dentro de dados/, fora do git.')
    plano=[]
    vistos=set()
    for r in propostas:
        if r.get('produto_id'): pid=str(int(r['produto_id']))
        else:
            handle=r['handle']
            encontrados=list(api.listar('products',handle=handle))
            if len(encontrados)!=1: raise ValueError('Link sem produto único; interrompido.')
            pid=str(int(encontrados[0]['id']))
        if pid in vistos: raise ValueError('Produto duplicado depois de resolver o link.')
        vistos.add(pid)
        for v in api._get(api._base+'/products/'+pid+'/variants').json():
            antes={k:v.get(k) for k in CAMPOS}
            plano.append({'produto_id':pid,'variante_id':v['id'],'antes':antes,'proposta':r['proposta'],
                          'depois':None,'estado':'simulação'})
    caminho.parent.mkdir(parents=True,exist_ok=True)
    def salvar(): caminho.write_text(json.dumps(plano,ensure_ascii=False,indent=2),encoding='utf-8')
    salvar()  # relatório antes de qualquer PUT; se não puder gravar, nenhuma escrita acontece
    if aplicar:
        for r in plano:
            if set(r['proposta'])!=set(CAMPOS): raise ValueError('Campo fora de peso/medidas bloqueado.')
            api._esperar_vaga()
            try:
                resposta=api._sessao.put(api._base+f"/products/{r['produto_id']}/variants/{r['variante_id']}",
                    headers=api._cabecalhos,json=r['proposta'],timeout=60,allow_redirects=False)
                api._guardar_limites(resposta)
                if resposta.status_code!=200: raise ValueError()
                retorno=resposta.json(); r['depois']={k:retorno.get(k) for k in CAMPOS}; r['estado']='aplicado'
            except Exception:
                r['estado']='não confirmado; conferir na loja'; salvar()
                raise ValueError('PUT não confirmado; interrompido, conferir relatório e loja.') from None
            salvar()
    return plano


@comando_seguro
def main():
    p=ArgumentParserSeguro(description='Simular correções de frete; aplicar só com autorização escrita.')
    p.add_argument('planilha'); p.add_argument('--aplicar',action='store_true'); p.add_argument('--autorizacao',default='')
    p.add_argument('--relatorio',default='dados/correcoes-frete/'+datetime.now().strftime('%Y%m%d-%H%M%S')+'.json')
    args=p.parse_args()
    if args.aplicar and args.autorizacao!=FRASE: p.error('Falta autorização escrita: '+FRASE)
    executar(NuvemshopAPI(carregar()),ler_propostas(args.planilha),args.relatorio,args.aplicar,args.autorizacao)
    print('Relatório gravado; '+('aplicação concluída.' if args.aplicar else 'simulação, nada alterado.'))


if __name__=='__main__': sys.exit(main())
